import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_heading_styled(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(27, 54, 93)  # Navy Blue
        p.paragraph_format.space_before = Pt(18)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(41, 128, 185) # Blue
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(52, 73, 94)  # Slate
    return p

def build_report_of_files():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base Normal Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Malgun Gothic'
    font.size = Pt(10)
    font.color.rgb = RGBColor(51, 51, 51)

    # ---------------------------------------------------------
    # COVER / TITLE SECTION
    # ---------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title_p.add_run("AUTH_ Anomaly_Detector - Master File Catalog & Technical Codebase Report")
    r.font.size = Pt(20)
    r.bold = True
    r.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Comprehensive Technical Documentation of All Subsystems, Source Files, Database Schemas & Data Flows")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(127, 140, 141)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # 1. EXECUTIVE SUMMARY & SYSTEM OVERVIEW
    # ---------------------------------------------------------
    add_heading_styled(doc, "1. Executive Summary & System Overview", level=1)
    
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "The AUTH_ Anomaly_Detector project is an enterprise-grade Single Sign-On (SSO) authentication security system designed "
        "to detect, classify, and neutralize cyber attacks in real time using Machine Learning. The codebase comprises a complete "
        "end-to-end architecture featuring a Random Forest Anomaly Detector, 2FA Email OTP Account Recovery Engine, Cryptographic "
        "Salted Password Hashing (scrypt/PBKDF2), Sliding-Window Rate Limiting, Production Cookie Hardening, and a standalone "
        "Red-Team Attacker Control Center for live attack simulation.\n\n"
        "This master document provides a comprehensive inventory and deep technical description of all 35+ source files, databases, "
        "frontend templates, utility scripts, and configuration manifests contained within the project repository."
    )

    # ---------------------------------------------------------
    # 2. MASTER FILE CATALOG & INVENTORY TABLE
    # ---------------------------------------------------------
    add_heading_styled(doc, "2. Master File Catalog & Inventory Table", level=1)

    p = doc.add_paragraph()
    p.add_run("The inventory table below details every file in the repository alongside its exact file size, line count, subsystem category, and core functional purpose.")

    file_inventory = [
        ("app.py", "53.4 KB", "1,269", "Core SSO & Backend Engine", "Main Flask WSGI application (Port 5000), ML Engine, Rate Limiter, 2FA Recovery & REST API endpoints."),
        (".env", "0.76 KB", "25", "Production Security Config", "Environment variable configuration file for Secret Key, Cookie security flags, Rate limit thresholds, and SMTP credentials."),
        (".env.example", "0.95 KB", "30", "Config Template", "Template configuration manifest for production deployment setup."),
        ("run_system.py", "2.37 KB", "65", "System Launcher", "Master CLI startup script for initializing databases and launching the main SSO Flask web server."),
        ("generate_users.py", "5.73 KB", "147", "Data Generation Utility", "Generates 200 student accounts pre-seeded with Werkzeug scrypt salted password hashes into users.db, users.json, and users.csv."),
        ("extract_features.py", "7.24 KB", "186", "ML Feature Pipeline", "Extracts 6D temporal & behavioral feature vectors from authentication logs using a 5-minute sliding window."),
        ("train_eval_models.py", "4.71 KB", "128", "ML Model Benchmark", "Trains and benchmarks 4 ML classifiers (Random Forest, Decision Tree, Logistic Regression, SVM) and selects the best model."),
        ("live_attack_simulator.py", "16.6 KB", "372", "Attack Vector Simulator", "Modular red-team attack simulator executing 4 live attack vectors and a continuous background traffic loop."),
        ("Attacker/attacker_app.py", "6.59 KB", "173", "Attacker Web Server", "Standalone Flask web server (Port 5001) serving the interactive Attacker Control Center Web Dashboard."),
        ("Attacker/templates/index.html", "21.9 KB", "470", "Attacker Console GUI", "Web GUI template for launching targeted cyber attack vectors, controlling continuous traffic, and monitoring live logs."),
        ("templates/index.html", "44.3 KB", "737", "SSO Frontend SPA", "Main SSO single-page web UI serving Admin Console, Security Analyst Console, Student Dashboard, 2FA Recovery Modal, and Real-Time Password Strength Visual Box."),
        ("static/js/app.js", "40.5 KB", "917", "Frontend Client JS", "Vanilla JS application logic managing Fetch API requests, 10-item pagination, 2FA modal handlers, and real-time password rule visual checklist."),
        ("static/css/style.css", "2.59 KB", "106", "Custom Styling Assets", "Custom CSS for glowing badges, status pills (ACTIVE, SUSPENDED, BLOCKED), custom scrollbars, and keyframe animations."),
        ("static/js/tailwind.3.4.1.js", "412 KB", "65", "Tailwind CSS Framework", "Offline standalone build of Tailwind CSS 3.4.1 framework ensuring full UI styling without external CDN dependency."),
        ("static/css/font-inter.css", "0.21 KB", "7", "Typography Loader", "Font face CSS definition for Inter typography font."),
        ("static/css/font-jetbrains.css", "0.24 KB", "7", "Typography Loader", "Font face CSS definition for JetBrains Mono code typography font."),
        ("generate_5000_logs.py", "6.61 KB", "182", "Synthetic Log Generator", "Generates 5,000 synthetic authentication logs representing normal logins and attack bursts into auth_logs.db."),
        ("unblock_cli.py", "1.41 KB", "37", "Admin CLI Tool", "Command-line utility allowing administrators to inspect blocked IPs, unblock IPs, and restore user account statuses."),
        ("check_db.py", "0.70 KB", "20", "DB Diagnostic Script", "Quick diagnostic tool checking SQLite database record counts and table schemas."),
        ("simulate_traffic.py", "5.02 KB", "115", "Traffic Generator", "Multi-threaded traffic simulator firing HTTP POST requests to test server concurrency and logging."),
        ("users.db", "69.6 KB", "476", "SQLite User Database", "Primary SQLite database storing 200 student accounts with scrypt salted password hashes, roles, and account statuses."),
        ("auth_logs.db", "1.03 MB", "1,555", "SQLite Event Database", "Primary security event audit trail database storing real-time authentication logs, risk scores, and attack types."),
        ("attacker_logs.db", "45.0 KB", "180", "SQLite Attacker Database", "Isolated database recording attack simulation events executed from the Attacker Control Center."),
        ("auth.log", "99.9 KB", "332", "JSON Lines Audit File", "Real-time structured JSON Lines log file mirroring all authentication events for SIEM parsing."),
        ("users.json", "77.7 KB", "1,802", "Synced JSON Data", "Formatted JSON representation of user accounts synced atomically from users.db."),
        ("users.csv", "18.8 KB", "201", "Synced CSV Data", "CSV export of user accounts synced atomically from users.db."),
        ("extracted_features_updated.csv", "120 KB", "6,556", "ML Benchmark Dataset", "Preprocessed 6D feature vector dataset used for training and evaluating machine learning models."),
        ("generate_updated_report_docx.py", "12.8 KB", "243", "Report Automation Script", "Automated Python script building report_updated.docx containing system updates."),
        ("generate_report_docx.py", "14.3 KB", "286", "Baseline Report Script", "Baseline Python script utilizing python-docx to build original report.docx."),
        ("report_updated.docx", "40.5 KB", "317", "Updated Summary Report", "Generated Word document summarizing high-level system features and production updates."),
        ("README.md", "9.59 KB", "196", "System Documentation", "Markdown documentation providing overview, setup instructions, API references, and CLI usage."),
        ("requirements.txt", "0.17 KB", "8", "Dependency Manifest", "Manifest specifying required Python libraries (Flask, scikit-learn, pandas, python-dotenv, python-docx, etc.)."),
        ("scratch/test_2fa_rules.py", "2.35 KB", "61", "Automated Unit Test", "Automated test verifying password strength policy, OTP generation, and 2FA recovery flow."),
        ("scratch/test_suspended_2fa.py", "1.73 KB", "46", "Automated Unit Test", "Automated test verifying account suspension handling, 2FA unsuspend, and salted hash DB updates."),
        ("scratch/test_rate_limiter.py", "1.07 KB", "29", "Automated Unit Test", "Automated test verifying environment loading, session cookie security flags, and Rate Limiter HTTP 429 throttling.")
    ]

    table = doc.add_table(rows=len(file_inventory) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["File Name", "Size", "Lines", "Subsystem Category", "Functional Description & Purpose"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1B365D")
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for row_idx, (fname, fsize, flines, fcat, fdesc) in enumerate(file_inventory, start=1):
        row_cells = table.rows[row_idx].cells
        row_cells[0].text = fname
        row_cells[1].text = fsize
        row_cells[2].text = flines
        row_cells[3].text = fcat
        row_cells[4].text = fdesc
        for c in row_cells:
            set_cell_margins(c)
        if row_idx % 2 == 0:
            for c in row_cells:
                set_cell_background(c, "F2F4F4")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # 3. CORE BACKEND & APPLICATION ENGINE (app.py)
    # ---------------------------------------------------------
    add_heading_styled(doc, "3. Core Backend & Application Engine (project/app.py)", level=1)

    p = doc.add_paragraph()
    p.add_run(
        "The central backend engine of the system is located in project/app.py (1,269 lines, 53.4 KB). It operates a Flask WSGI "
        "application on Port 5000, serving as the main SSO authentication server, ML anomaly evaluation engine, 2FA account recovery handler, "
        "and REST API provider for the web consoles."
    )

    app_functions = [
        ("init_db()", "Initializes SQLite databases (users.db and auth_logs.db). Creates tables and schema structures for accounts, logs, and blocked entities."),
        ("check_user_password(stored, input)", "Verifies user credentials against stored salted hashes using Werkzeug check_password_hash (scrypt/PBKDF2). Includes fallback for legacy plaintext passwords."),
        ("RateLimiter.is_rate_limited(key)", "In-memory sliding window rate limiter. Tracks IP request timestamps over a 60-second window to prevent DoS flooding on /login (15 req/min) and 2FA OTP requests (5 req/min)."),
        ("send_real_email_otp(to_email, otp_code, student_id)", "Handles real 2FA OTP email dispatch using TLS/SSL SMTP (smtplib). Gracefully falls back to console logging when SMTP credentials are not specified in .env."),
        ("extract_incoming_features(ip, student_id, user_agent)", "Computes real-time 6D feature vectors using a 300-second (5-minute) sliding window over recent logs in auth_logs.db."),
        ("classify_attack_type(features, is_bot)", "Categorizes detected anomaly patterns into 4 threat classes: Brute Force / Account DoS, Single-IP Password Spraying, IP Rotation Spraying, or Automated Bot Traffic."),
        ("login() [/login POST/GET]", "Primary SSO authentication route. Performs rate limiting check, account status evaluation (SUSPENDED/INACTIVE), IP lockout check, credential verification, ML risk score calculation, and audit logging."),
        ("request_2fa_otp() [/api/user/2fa/request-otp]", "2FA OTP request endpoint. Validates that the submitted Student ID and Email match registered DB records, generates a 5-minute valid 6-digit OTP, and enforces attacker IP isolation."),
        ("verify_unsuspend_2fa() [/api/user/2fa/verify-unsuspend]", "2FA verification & password reset endpoint. Validates OTP code, enforces 5-point password strength policy, updates user password hash, restores status to ACTIVE, and logs recovery event."),
        ("validate_password_strength(password)", "Server-side password policy validator. Enforces minimum 8 characters, uppercase (A-Z), lowercase (a-z), digit (0-9), and special character (!@#$%^&*)."),
        ("update_user_password_and_status_in_db()", "Updates user password salted hash and account status in SQLite users.db and triggers atomic sync to users.json and users.csv."),
        ("get_admin_logs() & get_analyst_analytics()", "REST API endpoints serving paginated audit logs (10 per page), security metrics, anomaly percentages, and attack distribution chart feeds for Admin and Analyst dashboards.")
    ]

    for fname, fdesc in app_functions:
        add_heading_styled(doc, fname, level=3)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(fdesc)

    # ---------------------------------------------------------
    # 4. MACHINE LEARNING & FEATURE EXTRACTION PIPELINE
    # ---------------------------------------------------------
    add_heading_styled(doc, "4. Machine Learning & Feature Extraction Pipeline", level=1)

    p = doc.add_paragraph()
    p.add_run(
        "The machine learning subsystem consists of project/extract_features.py (186 lines, 7.2 KB) and project/train_eval_models.py (128 lines, 4.7 KB). "
        "Together, they process raw authentication logs into structured 6D feature vectors and train classification models."
    )

    feat_items = [
        ("failed_attempts_5m", "Recent authentication failure count from the requesting IP address within the last 5 minutes (300 seconds)."),
        ("time_interval_sec", "Time gap (in seconds) between consecutive authentication attempts from the same IP address."),
        ("ip_changes_5m", "Number of distinct IP addresses attempting to authenticate against the target user account within 5 minutes."),
        ("global_failed_5m", "Total system-wide authentication failures across all accounts and IPs within the last 5 minutes."),
        ("global_unique_ips_5m", "Total count of unique IP addresses causing authentication failures system-wide within 5 minutes."),
        ("ua_bot_flag", "Binary indicator flag (0/1) identifying known automated bot User-Agents (e.g. Hydra, Python-urllib, Curl).")
    ]

    for name, desc in feat_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(f"{name}: ")
        r_b.bold = True
        bp.add_run(desc)

    add_heading_styled(doc, "ML Model Benchmark & Performance Evaluation Summary", level=2)

    ml_table = doc.add_table(rows=5, cols=7)
    ml_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    ml_headers = ["Classifier Model", "Accuracy (%)", "Precision (%)", "Recall (%)", "F1-Score (%)", "FPR (%)", "Latency (ms)"]
    hdr_cells_ml = ml_table.rows[0].cells
    for i, title in enumerate(ml_headers):
        hdr_cells_ml[i].text = title
        set_cell_background(hdr_cells_ml[i], "2980B9")
        hdr_cells_ml[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        hdr_cells_ml[i].paragraphs[0].runs[0].bold = True

    ml_data = [
        ("Logistic Regression", "95.98%", "97.46%", "91.88%", "94.59%", "1.48%", "0.00 ms"),
        ("Decision Tree", "96.44%", "97.89%", "92.68%", "95.21%", "1.23%", "0.00 ms"),
        ("Random Forest (Best Model)", "96.49%", "98.16%", "92.54%", "95.27%", "1.07%", "0.01 ms"),
        ("Support Vector Machine (SVM)", "95.83%", "97.05%", "91.88%", "94.39%", "1.73%", "0.05 ms")
    ]

    for row_idx, data in enumerate(ml_data, start=1):
        row_cells = ml_table.rows[row_idx].cells
        for col_idx, text in enumerate(data):
            row_cells[col_idx].text = text
            set_cell_margins(row_cells[col_idx])
            if row_idx == 3:
                set_cell_background(row_cells[col_idx], "EAFAF1")
                row_cells[col_idx].paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # 5. SECURITY MECHANISMS & PRODUCTION HARDENING
    # ---------------------------------------------------------
    add_heading_styled(doc, "5. Security Mechanisms & Production Hardening", level=1)

    sec_rules = [
        ("Cryptographic Salted Password Hashing", 
         "All user passwords in users.db, users.json, and users.csv are hashed using Werkzeug scrypt/PBKDF2 salted password hashing. "
         "This prevents password exposure in data leaks and neutralizes rainbow table attacks."),

        ("Real-Time Visual Password Strength Checklist", 
         "The frontend modal (index.html & app.js) includes a live 5-rule password checklist (#passwordRulesBox). As the user types, "
         "rule indicators dynamically update (✓ green / ✕ gray). Submission is blocked until all 5 rules (8+ chars, upper, lower, digit, special) pass."),

        ("2FA Email OTP Account Recovery Engine", 
         "Suspended users undergo 2FA Email OTP verification to unlock their account. Validates Student ID + Email match, generates a 6-digit 300s OTP, "
         "and blocks attacker IPs from requesting OTPs for third-party accounts."),

        ("Sliding-Window Rate Limiting (DoS Defense)", 
         "An in-memory RateLimiter protects endpoints against brute-force and HTTP flooding. /login is capped at 15 req/min per IP, returning HTTP 429 on breach."),

        ("Production Cookie & Session Security", 
         "Session configurations include SESSION_COOKIE_HTTPONLY = True, SESSION_COOKIE_SAMESITE = 'Lax', and a 30-minute expiration lifetime. "
         "Environment variables are managed dynamically via .env."),

        ("Immutable Admin Protection Mode", 
         "The primary administrator account 'admin' is protected against status modification, suspension, or deletion.")
    ]

    for title, desc in sec_rules:
        add_heading_styled(doc, title, level=3)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(desc)

    # ---------------------------------------------------------
    # 6. LIVE ATTACK SIMULATOR MODULE
    # ---------------------------------------------------------
    add_heading_styled(doc, "6. Live Attack Simulator Module (project/Attacker/)", level=1)

    p = doc.add_paragraph()
    p.add_run(
        "The Attacker module is a dedicated subsystem for red-team testing and demonstration. It includes Attacker/attacker_app.py (173 lines, 6.6 KB), "
        "Attacker/templates/index.html (470 lines, 21.9 KB), and project/live_attack_simulator.py (372 lines, 16.6 KB)."
    )

    attack_vectors = [
        ("Brute Force / Account DoS", "Fires 5 rapid failed login attempts from a single IP against a specific student account, causing status to flip to SUSPENDED and triggering IP lockout."),
        ("Single-IP Password Spraying", "Uses a single IP (198.51.100.42) to attempt a common password across multiple student accounts. Detected as Password Spraying by the ML engine."),
        ("IP Rotation Spraying", "Rotates across multiple public IPs while attacking different student accounts, testing multi-IP anomaly correlation."),
        ("Automated Bot Traffic", "Sends HTTP requests containing bot User-Agent strings (Hydra/v9.1, Python-urllib, Curl) to test instant User-Agent bot blocking."),
        ("Continuous Background Loop", "Launches a continuous background thread generating randomized real-time legitimate and malicious traffic bursts.")
    ]

    for title, desc in attack_vectors:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(f"{title}: ")
        r_b.bold = True
        bp.add_run(desc)

    # ---------------------------------------------------------
    # 7. WEB FRONTEND & UI ASSETS
    # ---------------------------------------------------------
    add_heading_styled(doc, "7. Web Frontend & UI Assets (project/templates/ & project/static/)", level=1)

    p = doc.add_paragraph()
    p.add_run(
        "The user interface is a modern Single-Page Application (SPA) built with HTML5, Tailwind CSS 3.4.1, and Vanilla JS:\n"
        "- project/templates/index.html (737 lines, 44.3 KB): SPA layout serving Admin Console, Analyst Console, Student Dashboard, 2FA Recovery Modal, and Password Strength UI Box.\n"
        "- project/static/js/app.js (917 lines, 40.5 KB): JS logic handling REST API calls, 10-item pagination, real-time password rule visual updates, and 2FA OTP flow.\n"
        "- project/static/css/style.css (106 lines, 2.59 KB): Custom animations, status pill badges (ACTIVE, SUSPENDED, BLOCKED), and scrollbars.\n"
        "- project/static/js/tailwind.3.4.1.js (412 KB): Offline build of Tailwind CSS framework."
    )

    # ---------------------------------------------------------
    # 8. DATABASES & STORAGE FORMATS
    # ---------------------------------------------------------
    add_heading_styled(doc, "8. Databases, Data Files & Storage Formats", level=1)

    db_items = [
        ("users.db (69.6 KB, SQLite)", "Primary directory database storing 200 student accounts with salted password hashes, roles, and statuses."),
        ("auth_logs.db (1.03 MB, SQLite)", "Primary security event log database recording authentication attempts, risk scores, and attack classifications."),
        ("attacker_logs.db (45.0 KB, SQLite)", "Isolated database recording attack simulation events executed from the Attacker Control Center."),
        ("auth.log (99.9 KB, JSON Lines)", "Real-time JSON Lines audit file logging authentication attempts for external SIEM integration."),
        ("users.json (77.7 KB) & users.csv (18.8 KB)", "Synced export data files mirroring users.db content for human readability and inspection.")
    ]

    for fname, fdesc in db_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)
        r_b = bp.add_run(f"{fname}: ")
        r_b.bold = True
        bp.add_run(fdesc)

    # ---------------------------------------------------------
    # 9. INTER-MODULE DEPENDENCY & DATA FLOW
    # ---------------------------------------------------------
    add_heading_styled(doc, "9. Inter-Module Dependency & Data Flow", level=1)

    p = doc.add_paragraph()
    p.add_run(
        "The end-to-end operational data flow across the codebase is structured as follows:\n\n"
        "1. Directory Seeding: generate_users.py initializes users.db with 200 student accounts using scrypt salted hashes, syncing to users.json and users.csv.\n"
        "2. Log Generation & ML Training: generate_5000_logs.py populates auth_logs.db. extract_features.py generates 6D feature vectors into extracted_features_updated.csv, and train_eval_models.py trains the Random Forest classifier.\n"
        "3. Runtime Authentication: app.py runs on Port 5000. When /login receives requests, RateLimiter checks frequency, extract_incoming_features() calculates 6D metrics, ML engine predicts Risk Score %, and security controls enforce IP lockout or account suspension.\n"
        "4. 2FA Account Recovery: Suspended users submit Student ID + Email on /api/user/2fa/request-otp. OTP is dispatched via SMTP/Console. Upon submission to /api/user/2fa/verify-unsuspend, password strength is checked, salted hash is saved, and status is restored to ACTIVE.\n"
        "5. Red-Team Attack Simulation: Attacker/attacker_app.py (Port 5001) triggers live attack payloads via live_attack_simulator.py against app.py, displaying real-time detection results on the Attacker Console and SSO Dashboard."
    )

    output_path_primary = "Reportoffiles_Updated.docx"
    doc.save(output_path_primary)
    print(f"[OK] Master report generated successfully: {os.path.abspath(output_path_primary)}")

    output_path_secondary = "Reportoffiles.docx"
    try:
        doc.save(output_path_secondary)
        print(f"[OK] Overwrote secondary file: {os.path.abspath(output_path_secondary)}")
    except PermissionError:
        print(f"[!] 'Reportoffiles.docx' is currently opened in Microsoft Word. Saved to '{output_path_primary}' instead.")

if __name__ == "__main__":
    build_report_of_files()
